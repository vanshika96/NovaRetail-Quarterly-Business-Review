from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
file_path ="novaretail_template.pptx"
prs=Presentation(file_path)

print("Number of slides",len(prs.slides))
print("Slide dimensions:",
            round(prs.slide_width/914400,2),"x",
            round(prs.slide_height/914400,2),"inches")

for slide_no,slide in enumerate(prs.slides,start=1):
    print(f"\n{'='*60}")
    print(f"SLIDE{slide_no}")
    print(f"{'='*60}")

    for shape in slide.shapes:
        text=""

        if shape.has_text_frame:
               text=shape.text.replace("\m","|")[:150]
               print(
                    f"Name:{shape.name!r} |"
                    f"Type:{shape.shape_type} |"
                    f"Position: ({shape.left / 914400:.2f},"
                    f"{shape.top / 914400:.2f}) |"
                    f"Size:({shape.width /914400:.2f},"
                    f"{shape.height / 914400:.2f})|"
                    f"text:{text!r}"
                                                                                                                           
               )
