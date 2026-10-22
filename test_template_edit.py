from pptx import Presentation
INPUT_FILE = "novaretail_template.pptx"
OUTPUT_FILE="template_test_output.pptx"

prs=Presentation(INPUT_FILE)

def find_shape(slide,shape_name):
    for shape in slide.shapes:
        if  shape.name== shape_name:
            return shape
    raise ValueError(f"Shape not found:{shape_name}")

#Test updating 1 placeholder

slide1 =prs.slides[0] 

title=find_shape(slide1,"deck_title")
title.text="NovaRetail - Quarterly Business Review"

#Update 1st KPI
kpi = find_shape(slide1,"kpi1_value")
kpi.text = "$1,280M"

for shape in slide1.shapes:
    print(repr(shape.name))
label = find_shape(slide1,"kpi1_label")
label.text ="Q2 2026 Revenue"

prs.save(OUTPUT_FILE)

print(f"Saved test ppt:{OUTPUT_FILE}")


